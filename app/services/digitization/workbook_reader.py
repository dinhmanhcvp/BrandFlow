"""
Đọc workbook thành lưới ô (grid) giữ nguyên thông tin mà pandas làm mất:
 • ô merge (vùng + ô gốc), để header nhiều tầng không biến thành 'Unnamed'
 • công thức (đọc lần 2 với data_only=False) — logic kinh doanh nằm trong công thức
 • vị trí ô thật trên Excel, để mọi con số truy ngược được (lineage)
"""

import os
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Dict, List, Optional, Tuple

from app.services.digitization.normalizers_vi import normalize_text


@dataclass
class SheetGrid:
  name: str
  cells: List[List[Any]]                      # cells[r][c], 0-based; chỉ ô gốc của vùng merge có giá trị
  merges: Dict[Tuple[int, int], Tuple[int, int, int, int]] = field(default_factory=dict)
  # (r, c) bất kỳ trong vùng merge → (r0, c0, r1, c1) của vùng đó (0-based, bao gồm)
  formulas: Dict[Tuple[int, int], str] = field(default_factory=dict)

  @property
  def n_rows(self) -> int:
    return len(self.cells)

  @property
  def n_cols(self) -> int:
    return max((len(r) for r in self.cells), default=0)

  def value(self, r: int, c: int) -> Any:
    if r < len(self.cells) and c < len(self.cells[r]):
      return self.cells[r][c]
    return None

  def filled_value(self, r: int, c: int) -> Any:
    """Giá trị ô, lấy từ ô gốc nếu ô nằm trong vùng merge."""
    region = self.merges.get((r, c))
    if region:
      return self.value(region[0], region[1])
    return self.value(r, c)

  def merge_of(self, r: int, c: int) -> Optional[Tuple[int, int, int, int]]:
    return self.merges.get((r, c))


def _clean_cell(value: Any) -> Any:
  if value is None:
    return None
  if isinstance(value, str):
    return normalize_text(value)
  if isinstance(value, float) and value != value:
    return None
  if isinstance(value, time):
    return value.isoformat()
  return value


def _read_xlsx(path: str) -> List[SheetGrid]:
  import openpyxl

  wb_values = openpyxl.load_workbook(path, data_only=True)
  try:
    wb_formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
  except Exception:
    wb_formulas = None

  sheets: List[SheetGrid] = []
  for ws in wb_values.worksheets:
    cells = [[_clean_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    # Bỏ các dòng trống ở cuối để không phải quét lưới rỗng khổng lồ
    while cells and all(v is None for v in cells[-1]):
      cells.pop()
    grid = SheetGrid(name=str(ws.title), cells=cells)

    for rng in ws.merged_cells.ranges:
      r0, c0, r1, c1 = rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1
      for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
          grid.merges[(r, c)] = (r0, c0, r1, c1)

    if wb_formulas is not None and ws.title in wb_formulas.sheetnames:
      for r_idx, row in enumerate(wb_formulas[ws.title].iter_rows(values_only=True)):
        if r_idx >= len(cells):
          break
        for c_idx, v in enumerate(row):
          if isinstance(v, str) and v.startswith("="):
            grid.formulas[(r_idx, c_idx)] = v
    sheets.append(grid)

  if wb_formulas is not None:
    wb_formulas.close()
  return sheets


def _read_with_pandas(path: str, ext: str) -> List[SheetGrid]:
  """CSV / XLS (không có thông tin merge) → lưới thô, header=None để không đoán sai header."""
  import pandas as pd

  def to_cells(df) -> List[List[Any]]:
    return [[_clean_cell(v) for v in row] for row in df.astype(object).itertuples(index=False)]

  if ext == "csv":
    last_error: Optional[Exception] = None
    for enc in ("utf-8", "utf-8-sig", "cp1258", "cp1252", "latin-1"):
      try:
        df = pd.read_csv(path, header=None, dtype=object, encoding=enc, sep=None, engine="python")
        return [SheetGrid(name=os.path.splitext(os.path.basename(path))[0], cells=to_cells(df))]
      except Exception as e:  # thử encoding tiếp theo
        last_error = e
    raise ValueError(f"Không đọc được CSV: {last_error}")

  frames = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
  return [SheetGrid(name=str(name), cells=to_cells(df)) for name, df in frames.items()]


def read_workbook(path: str) -> List[SheetGrid]:
  ext = os.path.splitext(path)[-1].lstrip(".").lower()
  if ext in ("xlsx", "xlsm"):
    return _read_xlsx(path)
  if ext in ("xls", "csv"):
    return _read_with_pandas(path, ext)
  raise ValueError(f"Định dạng .{ext} không phải bảng tính.")


TABULAR_EXTENSIONS = {"xlsx", "xlsm", "xls", "csv"}


def to_excel_col(c: int) -> str:
  """0 → 'A', 27 → 'AB'."""
  name = ""
  c += 1
  while c:
    c, rem = divmod(c - 1, 26)
    name = chr(65 + rem) + name
  return name


def is_number(value: Any) -> bool:
  return isinstance(value, (int, float)) and not isinstance(value, bool)


def is_date(value: Any) -> bool:
  return isinstance(value, (datetime, date))
